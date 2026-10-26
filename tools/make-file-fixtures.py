# Files a student really uploads, good and broken, for tests/files.test.mjs and the browser break check.
#   examples/14-word-notice.docx          a Word notice: paragraphs + an exam timetable TABLE (python-docx)
#   tests/fixtures/scanned-notice.pdf     an image-only PDF (no text layer) -> the OCR path
#   tests/fixtures/password.pdf           an encrypted PDF                  -> a clear "remove the password" message
#   tests/fixtures/corrupt.pdf            random bytes named .pdf
#   tests/fixtures/truncated.docx         the Word notice cut in half (a broken download)
#   tests/fixtures/not-really.docx        a text file renamed to .docx
#   tests/fixtures/old-word.doc           the header of an old binary .doc
#   tests/fixtures/zip-bomb.docx          a 30 KB .docx whose text inflates past 20 MB
#   tests/fixtures/empty.png              0 bytes
#   tests/fixtures/blank.png              a white picture with nothing to read
#   tests/fixtures/huge.png               12000 x 9000 pixels (a 108-megapixel decode)
# Run from the repo root: python tools/make-file-fixtures.py   (needs python-docx, Pillow, reportlab)
import io
import os
import random
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Pt
from PIL import Image
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
EX = ROOT / 'examples'
FX = ROOT / 'tests' / 'fixtures'
FX.mkdir(parents=True, exist_ok=True)
random.seed(26)


def word_notice():
    doc = Document()
    doc.styles['Normal'].font.size = Pt(12)
    doc.add_heading('NOTICE', level=1)
    doc.add_paragraph('Department of Computer Science & Engineering\tDate: 26/10/2026')
    doc.add_paragraph('All B.Tech III year students are informed that the Mid Semester Examination (MST-2) will be held as per the schedule below. Reporting time is 15 minutes before the exam.')
    t = doc.add_table(rows=1, cols=4)
    t.style = 'Table Grid'
    for cell, text in zip(t.rows[0].cells, ['Date', 'Subject', 'Time', 'Room']):
        cell.text = text
    for row in [('10 Nov 2026', 'Operating Systems', '10:00 AM to 11:30 AM', 'Room 204'),
                ('11 Nov 2026', 'DBMS', '10:00 AM to 11:30 AM', 'Room 204'),
                ('12 Nov 2026', 'Computer Networks', '2:00 PM to 3:30 PM', 'Lab 3')]:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    doc.add_paragraph('')
    p = doc.add_paragraph('Lab records must be submitted by 5 PM on 7 Nov 2026 in the department office.')
    p.runs[0].bold = True
    doc.add_paragraph('Note: Carry your ID card. Mobile phones are not allowed in the exam hall.')
    doc.add_paragraph('HoD, CSE')
    out = EX / '14-word-notice.docx'
    doc.save(out)
    print('wrote', out, out.stat().st_size, 'bytes')
    return out


def pdf_from_image(path, image):
    # JPEG inside the PDF, as a phone scanner app makes it (and small enough to keep in the repo)
    jpg = io.BytesIO()
    image.save(jpg, 'JPEG', quality=72)
    jpg.seek(0)
    from reportlab.lib.utils import ImageReader
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    c.drawImage(ImageReader(jpg), 30, 120, width=w - 60, height=(w - 60) * image.height / image.width)
    c.showPage()
    c.save()


def broken(docx_path):
    photo = Image.open(EX / '07-photo-notice.png').convert('RGB')
    photo.thumbnail((1200, 1200))
    pdf_from_image(FX / 'scanned-notice.pdf', photo)
    c = canvas.Canvas(str(FX / 'password.pdf'), pagesize=A4, encrypt='secret123')
    c.drawString(72, 760, 'Internal Test 2 on 13 Nov 2026 at 10 AM in Room 204.')
    c.save()
    (FX / 'corrupt.pdf').write_bytes(b'%PDF-1.7\n' + bytes(random.getrandbits(8) for _ in range(4000)))
    data = docx_path.read_bytes()
    (FX / 'truncated.docx').write_bytes(data[: len(data) // 2])
    (FX / 'not-really.docx').write_text('Internal Test 2 on 13 Nov 2026 at 10 AM in Room 204.', encoding='utf-8')
    (FX / 'old-word.doc').write_bytes(bytes.fromhex('D0CF11E0A1B11AE1') + bytes(504))
    body = '<w:p><w:r><w:t>' + ' ' * (24 * 1024 * 1024) + '</w:t></w:r></w:p>'
    with zipfile.ZipFile(FX / 'zip-bomb.docx', 'w', zipfile.ZIP_DEFLATED, compresslevel=9) as z:
        z.writestr('[Content_Types].xml', '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        z.writestr('word/document.xml', f'<?xml version="1.0"?><w:document xmlns:w="w"><w:body>{body}</w:body></w:document>')
    (FX / 'empty.png').write_bytes(b'')
    Image.new('RGB', (900, 600), (255, 255, 255)).save(FX / 'blank.png')
    Image.new('L', (12000, 9000), 255).save(FX / 'huge.png', optimize=True)
    for f in sorted(FX.iterdir()):
        print('wrote', f.relative_to(ROOT), f.stat().st_size, 'bytes')


if __name__ == '__main__':
    broken(word_notice())
